"""Appendix B must not silently fall behind `Settings` (M39 follow-up).

The manual was last rebuilt on 5 August for M49. By 12 August, **53 of 91
settings were mentioned nowhere in it**, including `QAT_CORPORATE_ACTION_MODE` -
the switch that decides whether M39 adjusts a stop or only logs what it would
have done - and `QAT_UI_LEVEL`, which gates the Workbench deploy button and the
Blotter's bulk sign-off.

Nobody had done anything wrong. A reference appendix maintained by hand goes
stale the moment a setting is added, and the gap is invisible because the
document still looks complete. That is the same failure mode as every asserted
count this project has got wrong, and the same fix applies: derive it.

So every field on `Settings` is either documented in `manual_body.py` or listed
in `_NOT_IN_THE_MANUAL` with a reason. The allowlist is the point of the test
rather than a hole in it - it turns "nobody noticed" into "somebody decided",
and the reason is what a later reader argues with.
"""

from __future__ import annotations

import pathlib

import pytest

from qat.config import Settings

_MANUAL_BODY = pathlib.Path(__file__).resolve().parent.parent / "scripts" / "manual_body.py"

# Deliberately undocumented, with the reason. Two kinds only: knobs whose value
# an operator has no reason to change, and plumbing that belongs to a deployment
# rather than to a user.
_NOT_IN_THE_MANUAL = {
    "ai_context_max_chars": "An internal bound on prompt size. Changing it alters cost and "
    "truncation behaviour and nothing an operator reasons about.",
    "ai_cost_budget_calls_per_day": "Advisory-layer spend cap. Real, but the AI layer cannot "
    "place an order, so it is not a trading control.",
    "alpaca_poll_seconds": "Feed cadence. Documented as a concept in the market-data section; "
    "the number is a vendor rate-limit detail.",
    "anthropic_model": "Provider model id. Changes with the vendor's catalogue, not with how "
    "the application behaves.",
    "bar_interval_seconds": "Aggregation interval. Changing it changes what every indicator "
    "means, so it is deliberately not offered as an operator knob.",
    "blotter_max_rows": "A display cap on one table. Raising it costs render time and shows "
    "older orders; it changes nothing about what the application does.",
    "correlation_window_bars": "Inside the correlation rail, which the manual documents by "
    "behaviour rather than by window length.",
    "data_dir": "Where the files live. The storage-location section documents the path "
    "itself, which is what an operator needs.",
    "database_url": "Deployment plumbing for the non-default storage backend, which no "
    "trading path reads - the ledger and the journal are CSV files on disk.",
    "equity_poll_seconds": "How often equity is sampled. The curve and the drawdown rail "
    "read whatever it produces, so the cadence is a resolution choice, not a rule.",
    "fred_series": "Which macro series are fetched. Changing the set changes the regime "
    "engine's inputs and is not an operator decision.",
    "ibkr_client_id": "IBKR connection plumbing, and that adapter is unimplemented beyond "
    "the seam.",
    "ibkr_host": "IBKR connection plumbing. Alpaca cannot reach the ASX at all, so this "
    "adapter is the direction of travel rather than a live path.",
    "ibkr_port": "IBKR connection plumbing, unreachable for the same reason as the host: "
    "that adapter exists as a seam with nothing implemented behind it.",
    "ibkr_pricing_model": "IBKR-specific, and unreachable while that adapter is a seam with "
    "nothing implemented behind it.",
    "local_llm_base_url": "Local-provider endpoint. The AI section covers choosing a "
    "provider; the URL is a deployment detail with no trading effect.",
    "local_llm_model": "Local-provider model id. Which model answers changes the wording of "
    "advice that cannot place an order, so it is not a trading control.",
    "reconciliation_poll_seconds": "How often tracked quantities are compared against the "
    "broker. The comparison is documented; its frequency is not a control.",
    "session_poll_seconds": "How often market hours are re-checked. The session rules are "
    "documented; how often they are consulted is not an operator decision.",
    "storage_backend": "sqlite or timescale. A deployment choice with no trading effect - "
    "the trading record is CSV on disk either way.",
    "watchlist_curated_asx": "The curated list itself, which the universe section describes "
    "by how it is built rather than by enumerating it.",
    "watchlist_curated_us": "As above: the universe section documents how the list is built "
    "and filtered, which is what decides whether a symbol can be traded.",
    "watchlist_min_avg_volume": "Inside the universe filter. The manual documents that thin "
    "names are excluded and why, which is the part that affects a decision.",
    "yfinance_poll_seconds": "Vendor request cadence, bounded by that vendor's tolerance "
    "rather than by anything about trading.",
}


def _manual_text() -> str:
    return _MANUAL_BODY.read_text(encoding="utf-8")


def _is_documented(field: str, text: str) -> bool:
    """Mentioned by env-var name, or by field name in prose.

    `QAT_` prefix first because that is what Appendix B tabulates and what an
    operator actually types.
    """
    return f"QAT_{field.upper()}" in text or field in text


def _fields() -> list[str]:
    return sorted(Settings.model_fields)


def test_there_are_settings_to_check():
    """A guard that silently checks nothing is worse than no guard - the lesson
    M78 had to learn about M75's colour check."""
    assert len(_fields()) >= 80


@pytest.mark.parametrize("field", _fields())
def test_every_setting_is_documented_or_excused(field):
    if field in _NOT_IN_THE_MANUAL:
        pytest.skip(f"declared undocumented: {_NOT_IN_THE_MANUAL[field]}")

    assert _is_documented(field, _manual_text()), (
        f"{field} is a setting with no mention in the manual. Document it in "
        f"manual_body.py, or add it to _NOT_IN_THE_MANUAL with a reason."
    )


def test_the_allowlist_does_not_outlive_its_entries():
    """An allowlist naming a setting that no longer exists is how a guard quietly
    stops covering what it claims to."""
    assert set(_NOT_IN_THE_MANUAL) <= set(Settings.model_fields)


def test_every_excuse_gives_a_reason():
    """ "Internal" is not a reason. The reason is what a later reader argues with,
    and it is the only thing separating this allowlist from a rubber stamp."""
    for field, reason in _NOT_IN_THE_MANUAL.items():
        assert len(reason) > 40, f"{field} needs a real reason, not a label"


def test_the_settings_that_gate_real_controls_are_never_excusable():
    """Some settings change what the application will DO, and no future reader
    gets to move them onto the allowlist without deleting this test.

    `ui_level` gates the Workbench deploy button and the Blotter's bulk sign-off,
    and its default is `standard` while the live config is `professional` - so a
    config reset silently removes a control. `corporate_action_mode` decides
    whether a split adjusts a resting stop or is merely logged, and an
    unadjusted stop through a split cost this account $375.23.
    """
    load_bearing = (
        "ui_level",
        "corporate_action_mode",
        "execution_mode",
        "trading_mode",
        "allow_autonomous_live_trading",
    )
    text = _manual_text()

    for field in load_bearing:
        assert field not in _NOT_IN_THE_MANUAL, f"{field} must never be excused"
        assert _is_documented(field, text), f"{field} must be documented"


# --- cross-references (M39 follow-up) -----------------------------------------


def test_no_manual_cross_reference_dangles():
    """Every "Section N.M" must name a heading that exists.

    Inserting two safety subsections and one Settings subsection renumbered the
    sections after them, and that silently broke FOUR pre-existing references -
    three pointing at the storage section and one at code signing. Each still
    read as a valid instruction and each now pointed somewhere else.

    Checked against the built document rather than the source, because the
    numbers only exist once the headings are laid down in order.
    """
    import re

    from docx import Document

    manual = _MANUAL_BODY.parent.parent / "docs" / "Quant_Advisory_Terminal_User_Manual.docx"
    if not manual.exists():
        pytest.skip("manual not built in this checkout - run `invoke manual`")

    document = Document(str(manual))
    headings = [p.text.strip() for p in document.paragraphs if p.style.name.startswith("Heading")]
    numbers = {h.split()[0].rstrip(".") for h in headings if h and h[0].isdigit()}
    text = "\n".join(p.text for p in document.paragraphs) + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )

    referenced = sorted(set(re.findall(r"Section (\d+(?:\.\d+)?)", text)))
    dangling = [ref for ref in referenced if ref not in numbers]

    assert referenced, "no cross-references found at all - the check would be vacuous"
    assert dangling == [], (
        f"these cross-references name sections that do not exist: {dangling}. "
        f"Renumbering a section breaks every reference to the ones after it."
    )
